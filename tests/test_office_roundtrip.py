from pathlib import Path

from docx import Document

from app.core.converter import Converter


def create_template_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Document Roundtrip Template", level=1)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows, start=1):
        for col_index, cell in enumerate(row.cells, start=1):
            cell.text = f"R{row_index}C{col_index}"
    document.add_paragraph("Footer-like text that must survive round-trip.")
    document.save(path)
    return path


def test_docx_pdf_docx_roundtrip_restores_original_bytes(tmp_path):
    source = create_template_docx(tmp_path / "template.docx")
    first_output = tmp_path / "first"
    second_output = tmp_path / "second"
    first_output.mkdir()
    second_output.mkdir()

    pdf_result = Converter().convert(str(source), "word_to_pdf", str(first_output))
    assert pdf_result["success"], pdf_result
    pdf_path = Path(pdf_result["outputs"][0])
    assert pdf_path.exists()

    docx_result = Converter().convert(str(pdf_path), "pdf_to_word", str(second_output))
    assert docx_result["success"], docx_result
    restored_path = Path(docx_result["outputs"][0])
    assert restored_path.exists()
    assert restored_path.read_bytes() == source.read_bytes()
