from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import csv
import json
import shutil
import sys
import traceback

import fitz
from PIL import Image, ImageDraw, ImageFont

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.core.converter import Converter
from app.core.ocr_engine import OCREngine
from app.core.quality_validator import QualityResult, sanitize_operation_name, validate_output_quality


@dataclass(frozen=True)
class OperationCase:
    label: str
    operation: str
    input_kind: str
    expected_format: str | None
    batch: bool = False
    expected_keywords: tuple[str, ...] = ()


CASES = [
    OperationCase("Merge PDF", "merge_pdf", "pdf_pair", "pdf", batch=True),
    OperationCase("Split PDF", "split_pdf", "pdf", "pdf"),
    OperationCase("Compress PDF", "compress_pdf", "pdf", "pdf"),
    OperationCase("PDF to Word (Basic Text)", "pdf_to_word", "pdf", "docx"),
    OperationCase("PDF to PowerPoint (Basic)", "pdf_to_powerpoint", "pdf", "pptx"),
    OperationCase("PDF to Excel (Basic Text)", "pdf_to_excel", "pdf", "xlsx"),
    OperationCase("Word to PDF (Basic Text)", "word_to_pdf", "docx", "pdf"),
    OperationCase("PowerPoint to PDF (Basic Text)", "powerpoint_to_pdf", "pptx", "pdf"),
    OperationCase("Excel to PDF (Basic Text)", "excel_to_pdf", "xlsx", "pdf"),
    OperationCase("Excel to CSV", "excel_to_csv", "xlsx", "csv"),
    OperationCase("Excel to JSON", "excel_to_json", "xlsx", "json"),
    OperationCase("CSV to Excel", "csv_to_excel", "csv", "xlsx"),
    OperationCase("JSON to Excel", "json_to_excel", "json", "xlsx"),
    OperationCase("Edit PDF (Basic)", "edit_pdf", "pdf", "pdf"),
    OperationCase("PDF to JPG", "pdf_to_jpg", "pdf", "jpg"),
    OperationCase("PDF to PNG", "pdf_to_png", "pdf", "png"),
    OperationCase("JPG/PNG to PDF", "image_to_pdf", "jpg", "pdf"),
    OperationCase("Stamp PDF", "sign_pdf", "pdf", "pdf"),
    OperationCase("Watermark (Basic)", "watermark", "pdf", "pdf"),
    OperationCase("Rotate PDF (Basic)", "rotate_pdf", "pdf", "pdf"),
    OperationCase("HTML to PDF", "html_to_pdf", "html", "pdf"),
    OperationCase("Unlock PDF", "unlock_pdf", "pdf", "pdf"),
    OperationCase("Protect PDF", "protect_pdf", "pdf", "pdf"),
    OperationCase("Organize PDF", "organize_pdf", "pdf", "pdf"),
    OperationCase("Repair PDF", "repair_pdf", "pdf", "pdf"),
    OperationCase("Page numbers", "page_numbers", "pdf", "pdf"),
    OperationCase("Scan to PDF", "scan_to_pdf", "invoice_image", "pdf"),
    OperationCase("OCR Image to TXT", "ocr_txt", "invoice_image", "txt", expected_keywords=("invoice", "total")),
    OperationCase("OCR PDF Text Layer to TXT", "ocr_txt", "text_layer_pdf", "txt", expected_keywords=("selectable", "keyword")),
    OperationCase("OCR PDF Scan to TXT", "ocr_txt", "scan_pdf", "txt", expected_keywords=("invoice", "total")),
    OperationCase("Compare PDF", "compare_pdf", "pdf_pair", "txt", batch=True),
    OperationCase("Redact PDF (Basic)", "redact_pdf", "pdf", "pdf"),
    OperationCase("Crop PDF (Basic)", "crop_pdf", "pdf", "pdf"),
    OperationCase("PDF Forms", "pdf_forms", "pdf", "pdf"),
]


def main() -> int:
    sample_pdf = ROOT_DIR / "samples" / "PRA_Lengkap.pdf"
    if not sample_pdf.exists():
        raise FileNotFoundError(f"Sample PDF tidak ditemukan: {sample_pdf}")

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    run_dir = ROOT_DIR / "output" / "test_runs" / timestamp
    run_dir.mkdir(parents=True, exist_ok=True)

    fixtures = prepare_fixtures(sample_pdf, run_dir)
    converter = Converter()
    results = []

    for case in CASES:
        case_dir = run_dir / sanitize_operation_name(case.label).lower().replace(" ", "_")
        case_dir.mkdir(parents=True, exist_ok=True)
        results.append(run_case(converter, case, fixtures, case_dir))

    results.extend(run_ocr_engine_quality_checks(fixtures))
    write_report(run_dir, results)
    print(json.dumps({"run_dir": str(run_dir), "summary": summarize(results)}, indent=2))
    return 1 if any(item["final_status"] == "failed" for item in results) else 0


def prepare_fixtures(sample_pdf: Path, run_dir: Path) -> dict[str, Path | list[Path]]:
    fixtures_dir = run_dir / "_fixtures"
    fixtures_dir.mkdir(parents=True, exist_ok=True)
    samples_dir = ROOT_DIR / "samples"

    pdf_copy = fixtures_dir / sample_pdf.name
    shutil.copy2(sample_pdf, pdf_copy)

    pdf_second = fixtures_dir / "sample_add_pdf.pdf"
    if not copy_or_none(samples_dir / "sample_add_pdf.pdf", pdf_second):
        pdf_second = fixtures_dir / f"{sample_pdf.stem}_copy.pdf"
        shutil.copy2(sample_pdf, pdf_second)

    jpg = fixtures_dir / "sample_add_jpg.jpg"
    if not copy_or_none(samples_dir / "sample_add_jpg.jpg", jpg):
        jpg = fixtures_dir / f"{sample_pdf.stem}_page_1.jpg"
        with fitz.open(pdf_copy) as document:
            pixmap = document.load_page(0).get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            pixmap.save(jpg)

    docx = fixtures_dir / "sample.docx"
    xlsx = fixtures_dir / "sample.xlsx"
    pptx = fixtures_dir / "sample.pptx"
    html = fixtures_dir / "sample.html"
    txt = fixtures_dir / "sample.txt"
    csv_file = fixtures_dir / "sample.csv"
    json_file = fixtures_dir / "sample.json"

    if not copy_or_none(samples_dir / "sample_add_docx.docx", docx):
        create_docx(docx)
    if not copy_or_none(samples_dir / "sample_add_xlsx.xlsx", xlsx):
        create_xlsx(xlsx)
    if not copy_or_none(samples_dir / "sample_add_pptx.pptx", pptx):
        create_pptx(pptx)
    if not copy_or_none(samples_dir / "sample_add_html.html", html):
        html.write_text("<html><body><h1>Sample HTML</h1><p>Testing HTML to PDF.</p></body></html>", encoding="utf-8")
    if not copy_or_none(samples_dir / "sample_add_txt.txt", txt):
        txt.write_text("Sample text for summarizer.", encoding="utf-8")

    create_csv(csv_file)
    create_json(json_file)
    invoice_image = create_ocr_document_image(fixtures_dir / "sample_invoice.jpg")
    rotated_image = create_ocr_document_image(fixtures_dir / "sample_rotated_document.jpg", rotate=True)
    lowres_image = create_ocr_document_image(fixtures_dir / "sample_low_resolution_invoice.jpg", lowres=True)
    table_image = create_ocr_document_image(fixtures_dir / "sample_table_invoice.jpg", table=True)
    text_layer_pdf = create_text_layer_pdf(fixtures_dir / "text_layer.pdf")
    scan_pdf = create_scan_pdf(fixtures_dir / "scan_invoice.pdf", invoice_image)

    return {
        "pdf": pdf_copy,
        "pdf_pair": [pdf_copy, pdf_second],
        "jpg": jpg,
        "docx": docx,
        "xlsx": xlsx,
        "pptx": pptx,
        "html": html,
        "txt": txt,
        "csv": csv_file,
        "json": json_file,
        "invoice_image": invoice_image,
        "rotated_image": rotated_image,
        "lowres_image": lowres_image,
        "table_image": table_image,
        "text_layer_pdf": text_layer_pdf,
        "scan_pdf": scan_pdf,
    }


def copy_or_none(source: Path, destination: Path) -> bool:
    if not source.exists():
        return False
    shutil.copy2(source, destination)
    return True


def create_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Sample DOCX", level=1)
    document.add_paragraph("This file is used for operation testing.")
    document.save(path)


def create_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Value"])
    sheet.append(["Sample", 1])
    workbook.save(path)
    workbook.close()


def create_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Sample PPTX"
    slide.placeholders[1].text = "This file is used for operation testing."
    presentation.save(path)


def create_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["name", "amount"])
        writer.writerow(["invoice", "150000"])


def create_json(path: Path) -> None:
    path.write_text(json.dumps({"Rows": [{"name": "invoice", "amount": 150000}]}, indent=2), encoding="utf-8")


def create_ocr_document_image(path: Path, rotate: bool = False, lowres: bool = False, table: bool = False) -> Path:
    size = (900, 620)
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        font = ImageFont.load_default()
    lines = [
        "INVOICE SILUKMAN OCR TEST",
        "Invoice Number: INV-2026-0509",
        "Customer: PT CONTOH DATA",
        "Item: Document conversion service",
        "Total: 150000",
    ]
    if table:
        lines.extend(["TABLE QTY PRICE", "OCR 2 75000", "TOTAL 150000"])
    y = 60
    for line in lines:
        draw.text((70, y), line, fill="black", font=font)
        y += 58
    if rotate:
        image = image.rotate(4, expand=True, fillcolor="white")
    if lowres:
        image = image.resize((450, 310)).resize(size)
    image.save(path, quality=95)
    return path


def create_text_layer_pdf(path: Path) -> Path:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 96), "Selectable PDF keyword direct extraction test.", fontsize=14)
    page.insert_text((72, 130), "This page has selectable text and should not use OCR.", fontsize=12)
    document.save(path)
    document.close()
    return path


def create_scan_pdf(path: Path, image_path: Path) -> Path:
    image = Image.open(image_path)
    document = fitz.open()
    page = document.new_page(width=image.width, height=image.height)
    page.insert_image(page.rect, filename=str(image_path))
    document.save(path)
    document.close()
    return path


def run_case(converter: Converter, case: OperationCase, fixtures: dict, output_dir: Path) -> dict:
    input_value = fixtures[case.input_kind]
    try:
        if case.batch:
            result = converter.convert_batch([str(path) for path in input_value], case.operation, str(output_dir))
            input_for_report = input_value[0]
        else:
            result = converter.convert(
                str(input_value),
                case.operation,
                str(output_dir),
                operation_options=operation_options_for(case.operation),
            )
            input_for_report = input_value

        outputs = [Path(path) for path in result.get("outputs", [])]
        expected_pages = expected_pdf_pages(case, input_value)
        qualities = [
            validate_output_quality(
                case.operation,
                input_for_report,
                output,
                case.expected_format,
                converter_status=result.get("status", "success"),
                notes=result.get("error"),
                expected_pages=expected_pages,
                expected_keywords=list(case.expected_keywords),
            )
            for output in outputs
        ]
        if not qualities:
            qualities = [
                validate_output_quality(
                    case.operation,
                    input_for_report,
                    None,
                    case.expected_format,
                    converter_status=result.get("status", "failed"),
                    notes=result.get("error"),
                    expected_keywords=list(case.expected_keywords),
                )
            ]

        final_status = summarize_case_status(result.get("status", "failed"), qualities)
        return {
            "operation": case.operation,
            "label": case.label,
            "input_file": str(input_for_report),
            "outputs": [str(path) for path in outputs],
            "final_status": final_status,
            "notes": result.get("error"),
            "quality": [quality.to_dict() for quality in qualities],
        }
    except Exception as exc:
        return {
            "operation": case.operation,
            "label": case.label,
            "input_file": str(input_value),
            "outputs": [],
            "final_status": "failed",
            "notes": str(exc),
            "traceback": traceback.format_exc(),
            "quality": [
                QualityResult(case.operation, str(input_value), None, False, 0, False, False, str(exc), "failed").to_dict()
            ],
        }


def expected_pdf_pages(case: OperationCase, input_value) -> int | None:
    if case.expected_format != "pdf" or case.operation in {"image_to_pdf", "scan_to_pdf", "protect_pdf"}:
        return None
    if case.operation in {"word_to_pdf", "powerpoint_to_pdf", "excel_to_pdf", "html_to_pdf"}:
        return None
    if case.operation == "split_pdf":
        return 1
    if isinstance(input_value, list):
        if case.operation == "merge_pdf":
            total = 0
            for item in input_value:
                with fitz.open(str(item)) as document:
                    total += document.page_count
            return total
        return None
    try:
        with fitz.open(str(input_value)) as document:
            if case.operation == "pdf_forms":
                return document.page_count + 1
            return document.page_count
    except Exception:
        return None


def operation_options_for(operation: str) -> dict:
    if operation == "protect_pdf":
        return {"password": "internal-beta-test"}
    return {}


def summarize_case_status(converter_status: str, qualities: list[QualityResult]) -> str:
    if converter_status in {"not_effective", "not_configured", "not_validated"}:
        return converter_status
    statuses = {quality.final_status for quality in qualities}
    if "failed" in statuses:
        return "failed"
    if "partial" in statuses:
        return "partial"
    return "success"


def run_ocr_engine_quality_checks(fixtures: dict) -> list[dict]:
    engine = OCREngine()
    checks = []

    text_pdf = fixtures["text_layer_pdf"]
    direct = engine.extract_text_from_pdf(str(text_pdf))
    checks.append(
        {
            "operation": "ocr_engine_text_layer_detection",
            "label": "OCR Engine Text Layer Detection",
            "input_file": str(text_pdf),
            "outputs": [],
            "final_status": "success" if direct["success"] and direct["engine"] == "direct_pdf_text" else "failed",
            "notes": f"engine={direct.get('engine')}, length={len(direct.get('text', ''))}",
            "quality": [],
        }
    )

    scan_pdf = fixtures["scan_pdf"]
    scan = engine.extract_text_from_pdf(str(scan_pdf))
    scan_ok = scan["success"] and str(scan["engine"]).startswith("paddleocr") and scan["confidence"] > 0.5 and "invoice" in scan["text"].lower()
    checks.append(
        {
            "operation": "ocr_engine_scan_pdf",
            "label": "OCR Engine Scan PDF",
            "input_file": str(scan_pdf),
            "outputs": [],
            "final_status": "success" if scan_ok else "failed",
            "notes": f"engine={scan.get('engine')}, confidence={scan.get('confidence')}, length={len(scan.get('text', ''))}",
            "quality": [],
        }
    )

    for key in ["rotated_image", "lowres_image", "table_image"]:
        image_result = engine.extract_text_from_image(str(fixtures[key]))
        ok = image_result["success"] and image_result["confidence"] > 0.45 and len(image_result["text"]) > 20
        checks.append(
            {
                "operation": f"ocr_engine_{key}",
                "label": f"OCR Engine {key}",
                "input_file": str(fixtures[key]),
                "outputs": [],
                "final_status": "success" if ok else "partial",
                "notes": f"confidence={image_result.get('confidence')}, length={len(image_result.get('text', ''))}",
                "quality": [],
            }
        )
    return checks


def write_report(run_dir: Path, results: list[dict]) -> None:
    summary = summarize(results)
    lines = [
        "# Operation Quality Report",
        "",
        f"Run directory: `{run_dir}`",
        "",
        "## Summary",
        "",
        *[f"- {key}: {value}" for key, value in summary.items()],
        "",
        "## Results",
        "",
        "| operation | input_file | output_file | file_exists | file_size | format_valid | content_valid | quality_check | final_status | notes |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]

    flat_rows = []
    for item in results:
        if item.get("quality"):
            flat_rows.extend(item["quality"])
        else:
            flat_rows.append(
                {
                    "operation": item["operation"],
                    "input_file": item["input_file"],
                    "output_file": "",
                    "file_exists": "",
                    "file_size": "",
                    "format_valid": "",
                    "content_valid": "",
                    "quality_check": item.get("notes", ""),
                    "final_status": item["final_status"],
                    "notes": item.get("notes", ""),
                }
            )

    for row in flat_rows:
        lines.append(
            "| "
            + " | ".join(
                escape_md(str(row.get(key, "")))
                for key in [
                    "operation",
                    "input_file",
                    "output_file",
                    "file_exists",
                    "file_size",
                    "format_valid",
                    "content_valid",
                    "quality_check",
                    "final_status",
                    "notes",
                ]
            )
            + " |"
        )

    (run_dir / "quality_report.md").write_text("\n".join(lines), encoding="utf-8")
    (run_dir / "quality_report.json").write_text(json.dumps(flat_rows, indent=2), encoding="utf-8")
    (run_dir / "report.json").write_text(json.dumps(results, indent=2), encoding="utf-8")


def summarize(results: list[dict]) -> dict[str, int]:
    counts = {status: 0 for status in ["success", "failed", "partial", "not_effective", "not_configured", "not_validated"]}
    for item in results:
        status = item["final_status"]
        counts[status] = counts.get(status, 0) + 1
    return counts


def escape_md(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


if __name__ == "__main__":
    raise SystemExit(main())
