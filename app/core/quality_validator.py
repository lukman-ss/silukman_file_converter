from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import csv
import json
import re
from zipfile import ZipFile

import fitz
from PIL import Image


FINAL_STATUSES = {
    "success",
    "success_short_text",
    "failed",
    "partial",
    "not_effective",
    "not_configured",
    "not_validated",
}


@dataclass
class QualityResult:
    operation: str
    input_file: str
    output_file: str | None
    file_exists: bool
    file_size: int
    format_valid: bool
    content_valid: bool
    quality_check: str
    final_status: str
    notes: str | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def sanitize_operation_name(name: str) -> str:
    invalid_chars = ['/', '\\', ':', '*', '?', '"', '<', '>', '|']
    for char in invalid_chars:
        name = name.replace(char, "_")
    return name


def validate_output_quality(
    operation: str,
    input_file: str | Path,
    output_file: str | Path | None,
    expected_format: str | None,
    converter_status: str = "success",
    notes: str | None = None,
    expected_pages: int | None = None,
    expected_keywords: list[str] | None = None,
    expected_ocr_pages: int | None = None,
    ocr_metadata: dict | None = None,
) -> QualityResult:
    input_path = Path(input_file)
    output_path = Path(output_file) if output_file else None
    exists = bool(output_path and output_path.exists())
    size = output_path.stat().st_size if exists and output_path else 0
    format_valid = False
    content_valid = False
    quality_check = ""
    final_status = converter_status if converter_status in FINAL_STATUSES else "failed"

    if final_status in {"not_effective", "not_configured", "not_validated"}:
        return QualityResult(
            operation,
            str(input_path),
            str(output_path) if output_path else None,
            exists,
            size,
            exists and size > 0,
            False,
            notes or final_status,
            final_status,
            notes,
        )

    if not output_path:
        return QualityResult(operation, str(input_path), None, False, 0, False, False, "No output file", "failed", notes)

    if not exists or size <= 0:
        return QualityResult(operation, str(input_path), str(output_path), exists, size, False, False, "Missing or empty output", "failed", notes)

    try:
        suffix = (expected_format or output_path.suffix.lstrip(".")).lower()
        if suffix == "pdf":
            format_valid, content_valid, quality_check = _validate_pdf(output_path, expected_pages)
        elif suffix in {"png", "jpg", "jpeg"}:
            format_valid, content_valid, quality_check = _validate_image(output_path)
        elif suffix == "docx":
            format_valid, content_valid, quality_check = _validate_docx(output_path)
        elif suffix == "xlsx":
            format_valid, content_valid, quality_check = _validate_xlsx(output_path)
        elif suffix == "pptx":
            format_valid, content_valid, quality_check = _validate_pptx(output_path)
        elif suffix == "txt":
            format_valid, content_valid, quality_check = _validate_text(
                output_path,
                expected_keywords,
                operation,
                expected_ocr_pages,
                ocr_metadata,
            )
        elif suffix == "csv":
            format_valid, content_valid, quality_check = _validate_csv(output_path)
        elif suffix == "json":
            format_valid, content_valid, quality_check = _validate_json(output_path)
        else:
            format_valid, content_valid, quality_check = True, True, "Basic non-empty validation"
    except Exception as exc:
        return QualityResult(operation, str(input_path), str(output_path), exists, size, False, False, str(exc), "failed", notes)

    if suffix == "txt" and quality_check.startswith("Short text accepted"):
        final_status = "success_short_text"
    elif converter_status == "partial":
        final_status = "partial"
    elif format_valid and content_valid:
        final_status = "success"
    elif format_valid:
        final_status = "partial"
    else:
        final_status = "failed"

    return QualityResult(
        operation,
        str(input_path),
        str(output_path),
        exists,
        size,
        format_valid,
        content_valid,
        quality_check,
        final_status,
        notes,
    )


def _validate_pdf(path: Path, expected_pages: int | None) -> tuple[bool, bool, str]:
    with fitz.open(path) as document:
        if document.is_encrypted:
            return True, True, "PDF encrypted as expected"
        page_count = document.page_count
        if page_count <= 0:
            return True, False, "PDF has no pages"
        if expected_pages is not None and page_count != expected_pages:
            return True, False, f"PDF page count {page_count}, expected {expected_pages}"
        return True, True, f"PDF page count {page_count}"


def _validate_image(path: Path) -> tuple[bool, bool, str]:
    with Image.open(path) as image:
        image.verify()
    with Image.open(path) as image:
        width, height = image.size
    return True, width > 0 and height > 0, f"Image readable {width}x{height}"


def _validate_docx(path: Path) -> tuple[bool, bool, str]:
    with ZipFile(path) as archive:
        names = archive.namelist()
        ok = "word/document.xml" in names
        media_count = len([name for name in names if name.startswith("word/media/")])
    return ok, ok, f"DOCX readable, embedded media {media_count}"


def _validate_xlsx(path: Path) -> tuple[bool, bool, str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = 0
        for sheet in workbook.worksheets:
            rows += sheet.max_row or 0
        return True, rows > 0, f"XLSX readable, sheets {len(workbook.sheetnames)}, rows {rows}"
    finally:
        workbook.close()


def _validate_pptx(path: Path) -> tuple[bool, bool, str]:
    from pptx import Presentation

    presentation = Presentation(path)
    slides = len(presentation.slides)
    return True, slides > 0, f"PPTX readable, slides {slides}"


def _validate_text(
    path: Path,
    expected_keywords: list[str] | None,
    operation: str,
    expected_ocr_pages: int | None,
    ocr_metadata: dict | None,
) -> tuple[bool, bool, str]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    text_length = len(text.strip())
    if expected_keywords:
        normalized_text = _normalize_text_for_keyword_match(text)
        missing = [
            keyword
            for keyword in expected_keywords
            if _normalize_text_for_keyword_match(keyword) not in normalized_text
        ]
        if missing:
            return True, False, f"Missing expected keywords: {', '.join(missing)}"

    ocr_checks = _validate_ocr_metadata(ocr_metadata, expected_ocr_pages) if operation == "ocr_txt" else []
    failing_ocr_checks = [item for item in ocr_checks if not item.startswith("ok:")]
    if failing_ocr_checks:
        return True, False, "; ".join(failing_ocr_checks)

    if text_length > 20:
        suffix = f"; {'; '.join(ocr_checks)}" if ocr_checks else ""
        return True, True, f"Text length {text_length}{suffix}"
    if text_length > 0 and operation == "ocr_txt":
        suffix = f"; {'; '.join(ocr_checks)}" if ocr_checks else ""
        return True, True, f"Short text accepted, length {text_length}{suffix}"
    return True, False, f"Text length {text_length}"


def _validate_ocr_metadata(ocr_metadata: dict | None, expected_ocr_pages: int | None) -> list[str]:
    if not ocr_metadata:
        return []

    checks = []
    confidence = float(ocr_metadata.get("confidence", 0) or 0)
    if confidence <= 0:
        checks.append("OCR confidence missing or zero")
    else:
        checks.append(f"ok: OCR confidence {confidence:.2f}")

    pages = ocr_metadata.get("pages") or []
    if expected_ocr_pages is not None:
        if len(pages) != expected_ocr_pages:
            checks.append(f"OCR page count {len(pages)}, expected {expected_ocr_pages}")
        else:
            checks.append(f"ok: OCR page count {len(pages)}")

    failed_pages = [
        page.get("page", index)
        for index, page in enumerate(pages, start=1)
        if page.get("status") != "success" or not str(page.get("text", "")).strip()
    ]
    if failed_pages:
        checks.append(f"OCR failed pages: {', '.join(str(page) for page in failed_pages)}")
    elif pages:
        checks.append("ok: all OCR pages have text")

    text = str(ocr_metadata.get("text", ""))
    if re.search(r"\b(?:Rp|IDR)\s*\d{1,3}(?:\.\d{3})+\b", text, flags=re.IGNORECASE):
        checks.append("ok: Indonesian currency format preserved")

    return checks


def _normalize_text_for_keyword_match(value: str) -> str:
    normalized = value.lower()
    normalized = re.sub(r"[\W_]+", " ", normalized, flags=re.UNICODE)
    return re.sub(r"\s+", " ", normalized).strip()


def _validate_csv(path: Path) -> tuple[bool, bool, str]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    return True, len(rows) > 0, f"CSV rows {len(rows)}"


def _validate_json(path: Path) -> tuple[bool, bool, str]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        count = len(data)
    elif isinstance(data, list):
        count = len(data)
    else:
        count = 1
    return True, count > 0, f"JSON top-level items {count}"
