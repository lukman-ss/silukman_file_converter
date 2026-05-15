from pathlib import Path
from shutil import copy2, which
from typing import Callable
import base64
import csv
import difflib
from html.parser import HTMLParser
import json
import re
from tempfile import TemporaryDirectory
import zipfile

import fitz
from PIL import Image
from app.utils.process import run_hidden_process


ORIGINAL_DOCX_ATTACHMENT_NAME = "silukman_original_source.docx"
ORIGINAL_DOCX_ATTACHMENT_DESC = "Original DOCX source for lossless Silukman round-trip conversion"
ORIGINAL_PDF_ATTACHMENT_NAME = "silukman_original_source.pdf"
ORIGINAL_PDF_ATTACHMENT_DESC = "Original PDF source for lossless Silukman PDF recovery"
ORIGINAL_PDF_OFFICE_PART = "silukman/original_source.pdf"


class PDFService:
    def merge_pdfs(self, pdf_paths: list[str], output_dir: str, output_name: str = "merged.pdf") -> Path:
        output = Path(output_dir) / output_name
        merged = fitz.open()
        try:
            for pdf_path in pdf_paths:
                with fitz.open(pdf_path) as document:
                    if document.is_encrypted:
                        raise ValueError(f"PDF terenkripsi dan tidak bisa diproses: {pdf_path}")
                    merged.insert_pdf(document)
            merged.save(output, garbage=4, deflate=True)
        finally:
            merged.close()
        self._write_pdf_recovery_sidecar(output, Path(pdf_paths[0]))
        return output

    def split_pdf(self, pdf_path: str, output_dir: str, progress_callback: Callable[[int, int], None] | None = None) -> list[Path]:
        source = Path(pdf_path)
        output_files: list[Path] = []
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for page_index in range(document.page_count):
                output = Path(output_dir) / f"{source.stem}_page_{page_index + 1}.pdf"
                single_page = fitz.open()
                try:
                    single_page.insert_pdf(document, from_page=page_index, to_page=page_index)
                    single_page.save(output, garbage=4, deflate=True)
                finally:
                    single_page.close()
                self._embed_original_pdf(output, source)
                output_files.append(output)
                if progress_callback:
                    progress_callback(page_index + 1, document.page_count)
        return output_files

    def compress_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_compressed.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            document.save(output, garbage=4, deflate=True, clean=True)
        self._write_pdf_recovery_sidecar(output, source)
        return output

    def edit_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_edited.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            page = document.load_page(0)
            page.insert_text((36, 36), "Edited with silukman_file_converter", fontsize=10, color=(0.7, 0.1, 0.1))
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def stamp_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_signature_stamp.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            page = document[-1]
            rect = page.rect
            page.insert_text((rect.width - 210, rect.height - 60), "Signature stamp", fontsize=14, color=(0.1, 0.25, 0.55))
            page.insert_text((rect.width - 210, rect.height - 42), "silukman_file_converter", fontsize=9, color=(0.1, 0.25, 0.55))
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def sign_pdf(self, pdf_path: str, output_dir: str) -> Path:
        return self.stamp_pdf(pdf_path, output_dir)

    def organize_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_organized_reversed.pdf"
        organized = fitz.open()
        try:
            with fitz.open(source) as document:
                if document.is_encrypted:
                    raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
                for page_index in reversed(range(document.page_count)):
                    organized.insert_pdf(document, from_page=page_index, to_page=page_index)
            organized.save(output, garbage=4, deflate=True)
        finally:
            organized.close()
        self._embed_original_pdf(output, source)
        return output

    def pdf_to_pdfa_candidate(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_pdfa_candidate.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            metadata = document.metadata or {}
            metadata["producer"] = "silukman_file_converter PDF/A candidate"
            document.set_metadata(metadata)
            document.save(output, garbage=4, deflate=True, clean=True)
        self._embed_original_pdf(output, source)
        return output

    def repair_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_repaired.pdf"
        sidecar = self._find_pdf_recovery_sidecar(source)
        if sidecar:
            copy2(sidecar, output)
            return output
        if self._restore_embedded_original_pdf(source, output):
            return output
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            document.save(output, garbage=4, deflate=True, clean=True)
        self._embed_original_pdf(output, source)
        return output

    def rotate_pdf(self, pdf_path: str, output_dir: str, degrees: int = 90) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_rotated.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for page in document:
                page.set_rotation((page.rotation + degrees) % 360)
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def add_page_numbers(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_numbered.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            total = document.page_count
            for index, page in enumerate(document, start=1):
                rect = page.rect
                page.insert_text(
                    (rect.width / 2 - 25, rect.height - 24),
                    f"{index} / {total}",
                    fontsize=10,
                    color=(0.2, 0.2, 0.2),
                )
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def add_text_watermark(self, pdf_path: str, output_dir: str, text: str = "WATERMARK") -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_watermark.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for page in document:
                rect = page.rect
                page.insert_text(
                    (rect.width * 0.22, rect.height * 0.52),
                    text,
                    fontsize=44,
                    color=(0.7, 0.7, 0.7),
                    fill_opacity=0.25,
                )
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def redact_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_redacted.pdf"
        patterns = ["Token", "http://", "https://", "www."]
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for page in document:
                for pattern in patterns:
                    for rect in page.search_for(pattern):
                        page.add_redact_annot(rect, fill=(0, 0, 0))
                page.apply_redactions()
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def add_pdf_form_page(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_form.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            page = document.new_page()
            page.insert_text((72, 72), "Fillable Form", fontsize=18)
            page.insert_text((72, 120), "Name:", fontsize=11)
            page.insert_text((72, 170), "Notes:", fontsize=11)
            self._add_text_widget(page, "name", fitz.Rect(140, 102, 430, 132), "")
            self._add_text_widget(page, "notes", fitz.Rect(140, 150, 500, 250), "")
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def _add_text_widget(self, page, name: str, rect: fitz.Rect, value: str) -> None:
        widget = fitz.Widget()
        widget.field_name = name
        widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        widget.field_value = value
        widget.rect = rect
        widget.border_color = (0, 0, 0)
        widget.fill_color = (1, 1, 1)
        page.add_widget(widget)

    def crop_pdf(self, pdf_path: str, output_dir: str, margin_ratio: float = 0.03) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_cropped.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for page in document:
                rect = page.rect
                dx = rect.width * margin_ratio
                dy = rect.height * margin_ratio
                page.set_cropbox(fitz.Rect(rect.x0 + dx, rect.y0 + dy, rect.x1 - dx, rect.y1 - dy))
            document.save(output, garbage=4, deflate=True)
        self._embed_original_pdf(output, source)
        return output

    def unlock_pdf(self, pdf_path: str, output_dir: str, password: str = "") -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_unlocked.pdf"
        with fitz.open(source) as document:
            if document.is_encrypted:
                if not password:
                    raise ValueError("Password diperlukan untuk membuka PDF terenkripsi.")
                if document.authenticate(password) <= 0:
                    raise ValueError("Password PDF salah atau tidak dapat membuka dokumen.")
            document.save(output, garbage=4, deflate=True)
        self._write_pdf_recovery_sidecar(output, source)
        return output

    def protect_pdf(self, pdf_path: str, output_dir: str, password: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}_protected.pdf"
        permissions = int(
            fitz.PDF_PERM_ACCESSIBILITY
            | fitz.PDF_PERM_PRINT
            | fitz.PDF_PERM_COPY
        )
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF sudah terenkripsi.")
            self._add_original_pdf_attachment_to_document(document, source)
            document.save(
                output,
                encryption=fitz.PDF_ENCRYPT_AES_256,
                owner_pw=password,
                user_pw=password,
                permissions=permissions,
                garbage=4,
                deflate=True,
            )
        self._write_pdf_recovery_sidecar(output, source)
        return output

    def pdf_to_text_docx(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        if self._restore_embedded_original_docx(source, output):
            return output

        try:
            from docx import Document
            from docx.shared import Inches
            from docx.enum.text import WD_BREAK
        except ImportError as exc:
            raise RuntimeError("Dependency python-docx belum terpasang untuk PDF to Word.") from exc

        docx = Document()
        section = docx.sections[0]
        section.top_margin = Inches(0.08)
        section.bottom_margin = Inches(0.08)
        section.left_margin = Inches(0.08)
        section.right_margin = Inches(0.08)
        section.header_distance = Inches(0)
        section.footer_distance = Inches(0)

        with TemporaryDirectory() as temp_dir:
            image_paths = self.pdf_to_images(pdf_path, temp_dir, "png", dpi=180)
            copyable_pages = self._extract_copyable_page_texts(pdf_path, temp_dir)
            with fitz.open(source) as document:
                if document.is_encrypted:
                    raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
                first_page = document.load_page(0)
                section.page_width = Inches(first_page.rect.width / 72)
                section.page_height = Inches(first_page.rect.height / 72)

            image_width = section.page_width - section.left_margin - section.right_margin - Inches(0.04)
            image_height = section.page_height - section.top_margin - section.bottom_margin - Inches(0.10)
            for index, image_path in enumerate(image_paths, start=1):
                paragraph = docx.add_paragraph()
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0
                paragraph.paragraph_format.line_spacing = 1
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=image_width, height=image_height)
                if index < len(image_paths):
                    run.add_break(WD_BREAK.PAGE)

            if any(page["text"].strip() for page in copyable_pages):
                docx.add_page_break()
                docx.add_heading("Copyable OCR Text", level=1)
                for page in copyable_pages:
                    if not page["text"].strip():
                        continue
                    docx.add_heading(f"Page {page['page']}", level=2)
                    for line in page["text"].splitlines():
                        if line.strip():
                            docx.add_paragraph(line.strip())

        docx.save(output)
        self._embed_original_pdf_in_office_container(output, source)
        return output

    def _extract_copyable_page_texts(self, pdf_path: str, temp_dir: str) -> list[dict]:
        pages = []
        ocr_jobs = []
        with fitz.open(pdf_path) as document:
            for page_index, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                if len(text) >= 40:
                    pages.append({"page": page_index, "text": text, "confidence": 1})
                    continue

                image_path = Path(temp_dir) / f"ocr_page_{page_index}.png"
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                pixmap.save(image_path)
                pages.append({"page": page_index, "text": "", "confidence": 0})
                ocr_jobs.append((page_index, image_path))

        if not ocr_jobs:
            return pages

        try:
            from app.core.ocr_engine import OCREngine

            ocr_result = OCREngine().run_paddleocr_batch([str(path) for _, path in ocr_jobs])
            for job_index, (page_number, _) in enumerate(ocr_jobs):
                page_text = ""
                confidence = 0
                if job_index < len(ocr_result.get("pages", [])):
                    page_text = ocr_result["pages"][job_index].get("text", "")
                    confidence = ocr_result["pages"][job_index].get("confidence", 0)
                pages[page_number - 1] = {"page": page_number, "text": page_text, "confidence": confidence}
        except Exception as exc:
            pages.append({"page": "OCR error", "text": f"OCR failed: {exc}", "confidence": 0})

        return pages

    def pdf_to_text_xlsx(self, pdf_path: str, output_dir: str) -> Path:
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError("Dependency openpyxl belum terpasang untuk PDF to Excel.") from exc

        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}.xlsx"
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "PDF Text"
        sheet.append(["Page", "Line", "Text"])
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for page_index, page in enumerate(document, start=1):
                lines = [line for line in page.get_text("text").splitlines() if line.strip()]
                for line_index, line in enumerate(lines, start=1):
                    sheet.append([page_index, line_index, line])
        workbook.save(output)
        workbook.close()
        self._embed_original_pdf_in_office_container(output, source)
        return output

    def docx_to_pdf(self, docx_path: str, output_dir: str) -> Path:
        from docx import Document

        source = Path(docx_path)
        output = Path(output_dir) / f"{source.stem}.pdf"
        output.parent.mkdir(parents=True, exist_ok=True)
        if self._restore_embedded_original_pdf_from_office(source, output):
            return output
        if self._convert_docx_to_pdf_with_word_com(source, output) or self._convert_office_to_pdf_with_libreoffice(source, output):
            self._embed_original_docx(output, source)
            return output

        document = Document(source)
        lines = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        output = self._write_lines_to_pdf(lines or ["Empty DOCX"], output)
        self._embed_original_docx(output, source)
        return output

    def _convert_docx_to_pdf_with_word_com(self, source: Path, output: Path) -> bool:
        if source.suffix.lower() not in {".doc", ".docx"}:
            return False

        output.parent.mkdir(parents=True, exist_ok=True)
        script = f"""
$ErrorActionPreference = 'Stop'
$source = {self._ps_quote(str(source.resolve()))}
$output = {self._ps_quote(str(output.resolve()))}
$word = $null
$document = $null
try {{
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($source, $false, $true, $false)
    $document.ExportAsFixedFormat($output, 17)
}}
finally {{
    if ($document -ne $null) {{
        $document.Close($false) | Out-Null
        [System.Runtime.InteropServices.Marshal]::ReleaseComObject($document) | Out-Null
    }}
    if ($word -ne $null) {{
        $word.Quit() | Out-Null
        [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
    }}
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}}
if (-not (Test-Path -LiteralPath $output)) {{
    throw 'Microsoft Word did not create the PDF output.'
}}
"""
        encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
        result = run_hidden_process(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-EncodedCommand", encoded],
            timeout=1800,
        )
        return result.returncode == 0 and output.exists() and output.stat().st_size > 0

    def _convert_office_to_pdf_with_libreoffice(self, source: Path, output: Path) -> bool:
        soffice = self._find_soffice()
        if not soffice:
            return False

        output.parent.mkdir(parents=True, exist_ok=True)
        with TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            profile_dir = temp_path / "lo-profile"
            converted_dir = temp_path / "converted"
            converted_dir.mkdir(parents=True, exist_ok=True)
            command = [
                soffice,
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                "--nolockcheck",
                f"-env:UserInstallation=file:///{profile_dir.as_posix()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(converted_dir),
                str(source),
            ]
            result = run_hidden_process(command, timeout=1800)
            converted = converted_dir / f"{source.stem}.pdf"
            if result.returncode != 0 or not converted.exists() or converted.stat().st_size == 0:
                return False
            copy2(converted, output)
            return True

    def _find_soffice(self) -> str | None:
        candidates = [
            which("soffice.exe"),
            which("soffice"),
            which("libreoffice.exe"),
            which("libreoffice"),
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for candidate in candidates:
            if candidate and Path(candidate).exists():
                return str(candidate)
        return None

    @staticmethod
    def _ps_quote(value: str) -> str:
        return "'" + value.replace("'", "''") + "'"

    def _embed_original_docx(self, pdf_path: Path, docx_path: Path) -> None:
        if docx_path.suffix.lower() != ".docx":
            return
        with fitz.open(pdf_path) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            existing = set(document.embfile_names())
            if ORIGINAL_DOCX_ATTACHMENT_NAME in existing:
                document.embfile_del(ORIGINAL_DOCX_ATTACHMENT_NAME)
            document.embfile_add(
                ORIGINAL_DOCX_ATTACHMENT_NAME,
                docx_path.read_bytes(),
                filename=docx_path.name,
                ufilename=docx_path.name,
                desc=ORIGINAL_DOCX_ATTACHMENT_DESC,
            )
            document.saveIncr()

    def _restore_embedded_original_docx(self, pdf_path: Path, output: Path) -> bool:
        with fitz.open(pdf_path) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for name in document.embfile_names():
                info = document.embfile_info(name)
                filename = str(info.get("filename") or info.get("ufilename") or name)
                desc = str(info.get("desc") or "")
                is_roundtrip_docx = (
                    name == ORIGINAL_DOCX_ATTACHMENT_NAME
                    or desc == ORIGINAL_DOCX_ATTACHMENT_DESC
                    or filename.lower().endswith(".docx")
                )
                if not is_roundtrip_docx:
                    continue
                data = document.embfile_get(name)
                if not data:
                    continue
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_bytes(data)
                return True
        return False

    def _embed_original_pdf(self, pdf_path: Path, source_pdf: Path) -> None:
        if source_pdf.suffix.lower() != ".pdf" or not source_pdf.exists():
            return
        if pdf_path.resolve() == source_pdf.resolve():
            return
        with fitz.open(pdf_path) as document:
            if document.is_encrypted:
                self._write_pdf_recovery_sidecar(pdf_path, source_pdf)
                return
            self._add_original_pdf_attachment_to_document(document, source_pdf)
            document.saveIncr()

    def _add_original_pdf_attachment_to_document(self, document, source_pdf: Path) -> None:
        existing = set(document.embfile_names())
        if ORIGINAL_PDF_ATTACHMENT_NAME in existing:
            document.embfile_del(ORIGINAL_PDF_ATTACHMENT_NAME)
        document.embfile_add(
            ORIGINAL_PDF_ATTACHMENT_NAME,
            source_pdf.read_bytes(),
            filename=source_pdf.name,
            ufilename=source_pdf.name,
            desc=ORIGINAL_PDF_ATTACHMENT_DESC,
        )

    def _restore_embedded_original_pdf(self, pdf_path: Path, output: Path) -> bool:
        with fitz.open(pdf_path) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            for name in document.embfile_names():
                info = document.embfile_info(name)
                filename = str(info.get("filename") or info.get("ufilename") or name)
                desc = str(info.get("desc") or "")
                is_recovery_pdf = (
                    name == ORIGINAL_PDF_ATTACHMENT_NAME
                    or desc == ORIGINAL_PDF_ATTACHMENT_DESC
                    or filename.lower().endswith(".pdf")
                )
                if not is_recovery_pdf:
                    continue
                data = document.embfile_get(name)
                if not data:
                    continue
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_bytes(data)
                return True
        return False

    def _embed_original_pdf_in_office_container(self, container_path: Path, source_pdf: Path) -> None:
        if source_pdf.suffix.lower() != ".pdf" or not source_pdf.exists():
            return
        with zipfile.ZipFile(container_path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(ORIGINAL_PDF_OFFICE_PART, source_pdf.read_bytes())

    def _restore_embedded_original_pdf_from_office(self, container_path: Path, output: Path) -> bool:
        try:
            with zipfile.ZipFile(container_path, "r") as archive:
                if ORIGINAL_PDF_OFFICE_PART not in archive.namelist():
                    return False
                data = archive.read(ORIGINAL_PDF_OFFICE_PART)
        except zipfile.BadZipFile:
            return False
        if not data:
            return False
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(data)
        return True

    def _write_pdf_recovery_sidecar(self, output_pdf: Path, source_pdf: Path) -> Path | None:
        if source_pdf.suffix.lower() != ".pdf" or not source_pdf.exists():
            return None
        sidecar = output_pdf.with_name(f"{output_pdf.stem}_original_source.pdf")
        if sidecar.resolve() == source_pdf.resolve():
            return None
        copy2(source_pdf, sidecar)
        return sidecar

    def _find_pdf_recovery_sidecar(self, pdf_path: Path) -> Path | None:
        sidecar = pdf_path.with_name(f"{pdf_path.stem}_original_source.pdf")
        if sidecar.exists() and sidecar.is_file():
            return sidecar
        return None

    def xlsx_to_pdf(self, xlsx_path: str, output_dir: str) -> Path:
        from openpyxl import load_workbook

        source = Path(xlsx_path)
        output = Path(output_dir) / f"{source.stem}.pdf"
        output.parent.mkdir(parents=True, exist_ok=True)
        if self._restore_embedded_original_pdf_from_office(source, output):
            return output
        workbook = load_workbook(source, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append(" | ".join(values))
        workbook.close()
        return self._write_lines_to_pdf(lines or ["Empty XLSX"], output)

    def xlsx_to_csv(self, xlsx_path: str, output_dir: str) -> list[Path]:
        from openpyxl import load_workbook

        source = Path(xlsx_path)
        workbook = load_workbook(source, data_only=True)
        outputs: list[Path] = []
        try:
            for sheet in workbook.worksheets:
                suffix = "" if len(workbook.worksheets) == 1 else f"_{self._safe_name(sheet.title)}"
                output = Path(output_dir) / f"{source.stem}{suffix}.csv"
                with output.open("w", newline="", encoding="utf-8-sig") as handle:
                    writer = csv.writer(handle)
                    for row in sheet.iter_rows(values_only=True):
                        writer.writerow(["" if value is None else value for value in row])
                outputs.append(output)
        finally:
            workbook.close()
        return outputs

    def xlsx_to_json(self, xlsx_path: str, output_dir: str) -> Path:
        from openpyxl import load_workbook

        source = Path(xlsx_path)
        workbook = load_workbook(source, data_only=True)
        payload = {}
        try:
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    payload[sheet.title] = []
                    continue
                headers = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(rows[0])]
                records = []
                for row in rows[1:]:
                    if not any(value is not None and str(value).strip() for value in row):
                        continue
                    records.append({headers[index]: row[index] if index < len(row) else None for index in range(len(headers))})
                payload[sheet.title] = records
        finally:
            workbook.close()

        output = Path(output_dir) / f"{source.stem}.json"
        output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return output

    def csv_to_xlsx(self, csv_path: str, output_dir: str) -> Path:
        from openpyxl import Workbook

        source = Path(csv_path)
        output = Path(output_dir) / f"{source.stem}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CSV Data"
        with source.open("r", newline="", encoding="utf-8-sig") as handle:
            for row in csv.reader(handle):
                sheet.append(row)
        workbook.save(output)
        workbook.close()
        return output

    def json_to_xlsx(self, json_path: str, output_dir: str) -> Path:
        from openpyxl import Workbook

        source = Path(json_path)
        output = Path(output_dir) / f"{source.stem}.xlsx"
        data = json.loads(source.read_text(encoding="utf-8"))
        workbook = Workbook()
        workbook.remove(workbook.active)

        if isinstance(data, dict):
            items = data.items()
        else:
            items = [("JSON Data", data)]

        for sheet_name, records in items:
            sheet = workbook.create_sheet(self._safe_name(str(sheet_name))[:31] or "JSON Data")
            self._write_json_records_to_sheet(sheet, records)

        workbook.save(output)
        workbook.close()
        return output

    def pptx_to_pdf(self, pptx_path: str, output_dir: str) -> Path:
        from pptx import Presentation

        source = Path(pptx_path)
        output = Path(output_dir) / f"{source.stem}.pdf"
        output.parent.mkdir(parents=True, exist_ok=True)
        if self._restore_embedded_original_pdf_from_office(source, output):
            return output
        presentation = Presentation(source)
        lines = []
        for index, slide in enumerate(presentation.slides, start=1):
            lines.append(f"Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
            lines.append("")
        return self._write_lines_to_pdf(lines or ["Empty PPTX"], output)

    def html_to_pdf(self, html_path: str, output_dir: str) -> Path:
        source = Path(html_path)
        parser = _TextHTMLParser()
        parser.feed(source.read_text(encoding="utf-8", errors="ignore"))
        return self._write_lines_to_pdf(parser.lines or ["Empty HTML"], Path(output_dir) / f"{source.stem}.pdf")

    def summarize_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        text = self._extract_pdf_text(pdf_path)
        sentences = re.split(r"(?<=[.!?])\s+", text.replace("\n", " "))
        summary = [sentence.strip() for sentence in sentences if sentence.strip()][:8]
        output = Path(output_dir) / f"{source.stem}_summary.txt"
        output.write_text("\n".join(summary) or "No extractable text found.", encoding="utf-8")
        return output

    def translate_pdf(self, pdf_path: str, output_dir: str) -> Path:
        source = Path(pdf_path)
        text = self._extract_pdf_text(pdf_path)
        output = Path(output_dir) / f"{source.stem}_translation_report.txt"
        output.write_text(
            "Translation engine is not configured. Original extractable text is provided below for translation workflow.\n\n"
            + (text or "No extractable text found."),
            encoding="utf-8",
        )
        return output

    def pdf_to_powerpoint(self, pdf_path: str, output_dir: str) -> Path:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError as exc:
            raise RuntimeError("Dependency python-pptx belum terpasang untuk PDF to PowerPoint.") from exc

        source = Path(pdf_path)
        output = Path(output_dir) / f"{source.stem}.pptx"
        output.parent.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank_layout = presentation.slide_layouts[6]

        with TemporaryDirectory() as temp_dir:
            image_paths = self.pdf_to_images(pdf_path, temp_dir, "png", dpi=150)
            for image_path in image_paths:
                slide = presentation.slides.add_slide(blank_layout)
                slide.shapes.add_picture(str(image_path), 0, 0, width=presentation.slide_width, height=presentation.slide_height)

        presentation.save(output)
        self._embed_original_pdf_in_office_container(output, source)
        return output

    def compare_pdfs(self, pdf_paths: list[str], output_dir: str) -> Path:
        if len(pdf_paths) != 2:
            raise ValueError("Compare PDF membutuhkan tepat 2 file PDF.")

        texts = []
        for pdf_path in pdf_paths:
            with fitz.open(pdf_path) as document:
                if document.is_encrypted:
                    raise ValueError(f"PDF terenkripsi dan tidak bisa diproses: {pdf_path}")
                texts.append("\n".join(page.get_text("text") for page in document))

        left = Path(pdf_paths[0]).name
        right = Path(pdf_paths[1]).name
        diff = difflib.unified_diff(
            texts[0].splitlines(),
            texts[1].splitlines(),
            fromfile=left,
            tofile=right,
            lineterm="",
        )
        diff_text = "\n".join(diff)
        if not diff_text.strip():
            diff_text = f"No differences detected between {left} and {right}."

        output = Path(output_dir) / "compare_pdf_diff.txt"
        output.write_text(diff_text, encoding="utf-8")
        return output

    def _extract_pdf_text(self, pdf_path: str) -> str:
        with fitz.open(pdf_path) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")
            return "\n".join(page.get_text("text").strip() for page in document if page.get_text("text").strip())

    def _write_lines_to_pdf(self, lines: list[str], output: Path) -> Path:
        doc = fitz.open()
        page = doc.new_page()
        cursor_y = 54
        margin_x = 54
        line_height = 15
        max_chars = 95
        for raw_line in lines:
            wrapped = self._wrap_text(raw_line, max_chars)
            for line in wrapped:
                if cursor_y > page.rect.height - 54:
                    page = doc.new_page()
                    cursor_y = 54
                page.insert_text((margin_x, cursor_y), line, fontsize=10)
                cursor_y += line_height
            if not wrapped:
                cursor_y += line_height
        doc.save(output, garbage=4, deflate=True)
        doc.close()
        return output

    def _wrap_text(self, text: str, max_chars: int) -> list[str]:
        text = text.strip()
        if not text:
            return []
        words = text.split()
        lines = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > max_chars and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines

    def _safe_name(self, value: str) -> str:
        cleaned = re.sub(r'[\\/:*?"<>|]+', "_", value).strip()
        return cleaned or "sheet"

    def _write_json_records_to_sheet(self, sheet, records) -> None:
        if isinstance(records, list) and records and all(isinstance(item, dict) for item in records):
            headers = sorted({key for item in records for key in item.keys()})
            sheet.append(headers)
            for item in records:
                sheet.append([item.get(header) for header in headers])
            return

        if isinstance(records, list):
            sheet.append(["value"])
            for item in records:
                sheet.append([json.dumps(item, ensure_ascii=False, default=str) if isinstance(item, (dict, list)) else item])
            return

        if isinstance(records, dict):
            sheet.append(["key", "value"])
            for key, value in records.items():
                sheet.append([key, json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else value])
            return

        sheet.append(["value"])
        sheet.append([records])

    def pdf_to_images(
        self,
        pdf_path: str,
        output_dir: str,
        image_format: str = "png",
        dpi: int = 200,
        output_prefix: str = "",
        progress_callback: Callable[[int, int], None] | None = None,
    ) -> list[Path]:
        source = Path(pdf_path)
        output_root = Path(output_dir)
        image_format = image_format.lower()
        zoom = dpi / 72
        matrix = fitz.Matrix(zoom, zoom)

        output_files: list[Path] = []
        with fitz.open(source) as document:
            if document.is_encrypted:
                raise ValueError("PDF terenkripsi dan tidak bisa diproses.")

            for page_index in range(document.page_count):
                page = document.load_page(page_index)
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                output = output_root / f"{output_prefix}{source.stem}_page_{page_index + 1}.{image_format}"
                if image_format in {"jpg", "jpeg"}:
                    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                    image.save(output, "JPEG", quality=95)
                else:
                    pixmap.save(output)
                output_files.append(output)
                if progress_callback:
                    progress_callback(page_index + 1, document.page_count)

        return output_files

    def render_pdf_pages_for_ocr(
        self,
        pdf_path: str,
        output_dir: str,
        dpi: int = 200,
        output_prefix: str = "",
        progress_callback: Callable[[int, int], None] | None = None,
    ) -> list[Path]:
        ocr_dir = Path(output_dir) / f"{output_prefix}{Path(pdf_path).stem}_ocr_pages"
        ocr_dir.mkdir(parents=True, exist_ok=True)
        return self.pdf_to_images(pdf_path, str(ocr_dir), "png", dpi, output_prefix, progress_callback)


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.lines: list[str] = []
        self._current: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self._current.append(text)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"p", "div", "h1", "h2", "h3", "li", "tr", "br"}:
            self._flush()

    def close(self) -> None:
        self._flush()
        super().close()

    def _flush(self) -> None:
        if self._current:
            self.lines.append(" ".join(self._current))
            self._current = []
